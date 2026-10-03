"""Word (.docx) and PowerPoint (.pptx) → elements."""
import re
from pathlib import Path

import docx
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .base import DocBuilder, save_media, table_markdown

MONO = re.compile(r"mono|courier|consol|menlo|cascadia|code", re.I)


def parse_docx(path: Path, title: str | None = None) -> DocBuilder:
    d = docx.Document(str(path))
    b = DocBuilder(title or d.core_properties.title or Path(path).stem)
    code_buf: list[str] = []

    def flush():
        if code_buf:
            b.add("code", "\n".join(code_buf))
            code_buf.clear()

    for child in d.element.body.iterchildren():
        tag = child.tag.split("}")[1]
        if tag == "tbl":
            flush()
            t = docx.table.Table(child, d)
            b.add("table", table_markdown([[c.text for c in r.cells] for r in t.rows]))
            continue
        if tag != "p":
            continue
        p = docx.text.paragraph.Paragraph(child, d)
        style = (p.style.name if p.style is not None else "") or ""
        text = p.text
        # images inside the paragraph
        for blip in child.iter(qn("a:blip")):
            rid = blip.get(qn("r:embed"))
            part = d.part.related_parts.get(rid)
            if part is not None and hasattr(part, "blob"):
                ext = Path(part.partname).suffix or ".png"
                flush()
                b.add("figure", "", media_path=save_media(part.blob, ext))
        if not text.strip():
            continue
        runs = [r for r in p.runs if r.text.strip()]
        is_code = "code" in style.lower() or (runs and all(r.font.name and MONO.search(r.font.name) for r in runs))
        if is_code:
            code_buf.append(text)
            continue
        flush()
        m = re.match(r"heading\s*(\d)", style, re.I)
        if m:
            b.heading(text, int(m.group(1)))
        elif style.lower() == "title":
            b.heading(text, 1)
        elif "list" in style.lower() or child.find(".//" + qn("w:numPr")) is not None:
            b.add("list_item", text)
        elif "quote" in style.lower():
            b.add("quote", text)
        elif "caption" in style.lower():
            b.add("caption", text)
        else:
            b.add("paragraph", text)
    flush()
    return b


def parse_pptx(path: Path, title: str | None = None) -> DocBuilder:
    prs = Presentation(str(path))
    b = DocBuilder(title or prs.core_properties.title or Path(path).stem)
    for n, slide in enumerate(prs.slides, 1):
        title_shape = slide.shapes.title
        stitle = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        b.heading(stitle or f"Slide {n}", 1, page=n)
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            _shape(b, shape, n)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                b.add("note", notes, page=n, meta={"speaker_notes": True})
    return b


def _shape(b: DocBuilder, shape, n: int):
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for s in shape.shapes:
            _shape(b, s, n)
        return
    if shape.has_table if hasattr(shape, "has_table") else False:
        rows = [[c.text for c in r.cells] for r in shape.table.rows]
        b.add("table", table_markdown(rows), page=n)
        return
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        try:
            b.add("figure", shape.name if not shape.name.lower().startswith("picture") else "", page=n,
                  media_path=save_media(shape.image.blob, shape.image.ext))
        except Exception:  # noqa: BLE001 - linked/missing images
            pass
        return
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs).strip()
            if t:
                b.add("list_item" if p.level > 0 or len(shape.text_frame.paragraphs) > 1 else "paragraph", t,
                      page=n, meta={"indent": p.level} if p.level else {})
