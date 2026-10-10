"""Spreadsheets (.xlsx, .csv, .tsv) and e-books (.epub) → elements."""
import csv
import io
import posixpath
import re
import zipfile
from pathlib import Path

from bs4 import BeautifulSoup

from .base import DocBuilder, table_markdown

MAX_ROWS = 400  # rows per table element; larger sheets are split into chunks


def _add_rows(b: DocBuilder, rows: list[list]):
    rows = [r for r in rows if any(str(c or "").strip() for c in r)]
    if not rows:
        return
    head, body = rows[0], rows[1:] or []
    for i in range(0, max(1, len(body)), MAX_ROWS):
        b.add("table", table_markdown([head] + body[i:i + MAX_ROWS]))


def parse_xlsx(path: Path, title: str | None = None) -> DocBuilder:
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    b = DocBuilder(title or Path(path).stem)
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if not any(any(c is not None for c in r) for r in rows):
            continue
        b.heading(ws.title, 1)
        _add_rows(b, rows)
    return b


def parse_csv(data: str, title: str, delimiter: str | None = None) -> DocBuilder:
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(data[:4000], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    b = DocBuilder(title)
    b.heading(title, 1)
    _add_rows(b, list(csv.reader(io.StringIO(data), delimiter=delimiter)))
    return b


def parse_epub(path: Path, title: str | None = None) -> DocBuilder:
    """EPUB is a zip of XHTML chapters; read them in spine order and reuse the HTML parser."""
    from .web import merge_pages, parse_html
    z = zipfile.ZipFile(str(path))
    container = BeautifulSoup(z.read("META-INF/container.xml"), "xml")
    opf_path = container.find("rootfile")["full-path"]
    opf = BeautifulSoup(z.read(opf_path), "xml")
    base = posixpath.dirname(opf_path)
    t = opf.find("dc:title") or opf.find("title")
    book_title = title or (t.get_text(strip=True) if t else "") or Path(path).stem
    items = {i["id"]: i["href"] for i in opf.find_all("item") if i.get("href")}
    order = [items[r["idref"]] for r in opf.find_all("itemref") if r.get("idref") in items]
    pages = []
    for href in order:
        name = posixpath.normpath(posixpath.join(base, href))
        if not re.search(r"\.x?html?$", name, re.I) or name not in z.namelist():
            continue
        html = z.read(name).decode("utf-8", "replace")
        pb = parse_html(html, f"epub:{name}", download_images=0)
        if sum(len(e.text) for e in pb.elements) > 200:  # skip covers, title pages, blank separators
            pages.append((name, pb))
    return merge_pages(book_title, pages)
