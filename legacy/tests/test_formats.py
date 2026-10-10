"""Spreadsheet and e-book parsers produce sections and tables without any model calls."""
import zipfile

from engine.ingest.sheets import parse_csv, parse_epub, parse_xlsx

PARA = "Photosynthesis turns light energy into chemical energy stored in glucose. " * 6


def test_csv_sniffs_delimiter_and_keeps_header():
    b = parse_csv("Element;Symbol;Group\nSodium;Na;1\nChlorine;Cl;17\n", "Elements")
    tables = [e for e in b.elements if e.kind == "table"]
    assert len(tables) == 1
    assert tables[0].text.splitlines()[0] == "| Element | Symbol | Group |"
    assert "| Chlorine | Cl | 17 |" in tables[0].text


def test_xlsx_one_section_per_sheet(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.title = "Prices"
    wb.active.append(["Item", "Price"])
    wb.active.append(["Tea", 20])
    ws = wb.create_sheet("Empty")
    ws2 = wb.create_sheet("Staff")
    ws2.append(["Name", "Role"])
    ws2.append(["Asha", "Lead"])
    f = tmp_path / "book.xlsx"
    wb.save(f)
    b = parse_xlsx(f)
    heads = [e.text for e in b.elements if e.kind == "heading"]
    assert heads == ["Prices", "Staff"] and ws.title == "Empty"
    assert sum(e.kind == "table" for e in b.elements) == 2


def test_epub_reads_chapters_in_spine_order(tmp_path):
    f = tmp_path / "book.epub"
    with zipfile.ZipFile(f, "w") as z:
        z.writestr("mimetype", "application/epub+zip")
        z.writestr("META-INF/container.xml", '<?xml version="1.0"?><container><rootfiles>'
                   '<rootfile full-path="OEBPS/content.opf"/></rootfiles></container>')
        z.writestr("OEBPS/content.opf", '<?xml version="1.0"?><package><metadata><dc:title>Plant Life</dc:title>'
                   '</metadata><manifest><item id="c1" href="ch1.xhtml"/><item id="c2" href="ch2.xhtml"/>'
                   '<item id="cov" href="cover.xhtml"/></manifest><spine><itemref idref="cov"/>'
                   '<itemref idref="c2"/><itemref idref="c1"/></spine></package>')
        z.writestr("OEBPS/cover.xhtml", "<html><body><p>Cover</p></body></html>")
        for n, t in (("ch1", "Roots"), ("ch2", "Leaves")):
            z.writestr(f"OEBPS/{n}.xhtml", f"<html><head><title>{t}</title></head><body><h1>{t}</h1><p>{PARA}</p></body></html>")
    b = parse_epub(f)
    assert b.title == "Plant Life"
    tops = [e.text for e in b.elements if e.kind == "heading" and e.level == 1]
    assert tops == ["Leaves", "Roots"]  # spine order, cover skipped
